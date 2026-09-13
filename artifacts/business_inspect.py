from pathlib import Path
import hashlib,json,re
from pypdf import PdfReader
from docx import Document
import pypdfium2 as pdfium

root=Path(r'C:\Users\user\Downloads\BD Business Docs')
out=Path('artifacts/business-inspection');out.mkdir(parents=True,exist_ok=True)
rows=[]
for i,p in enumerate(sorted(root.iterdir()),1):
    if not p.is_file(): continue
    row=dict(file=p.name,path=str(p),sha256=hashlib.sha256(p.read_bytes()).hexdigest(),bytes=p.stat().st_size,decision='pending inspection',metadata={},ids={},relationships=[],ocr_choice=None,status='local_only')
    if p.suffix.lower()=='.pdf':
        doc=PdfReader(p); pages=[x.extract_text() or '' for x in doc.pages]
        row.update(pages=len(pages),page_text_lengths=[len(x) for x in pages],pdf_metadata={str(k):str(v) for k,v in (doc.metadata or {}).items()})
        text='\n'.join(f'\n--- PAGE {n+1} ---\n{t}' for n,t in enumerate(pages))
        render=pdfium.PdfDocument(str(p))
        for n in sorted(set([0,len(pages)-1])):
            render[n].render(scale=1.6).to_pil().save(out/f'{i:02d}-page-{n+1}.png')
    elif p.suffix.lower()=='.docx':
        doc=Document(p);text='\n'.join(x.text for x in doc.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows)
    else:text=p.read_text(encoding='utf-8-sig')
    dest=out/f'{i:02d}.txt'; dest.write_text(text,encoding='utf-8');row['inspection_text']=str(dest);row['characters']=len(text)
    rows.append(row)
    print(json.dumps(dict(number=i,file=p.name,pages=row.get('pages'),characters=len(text),head=text[:1700],tail=text[-900:]),ensure_ascii=False))
Path('business-upload-log.json').write_text(json.dumps(dict(project_id='f932c9e0-40af-4ff3-8704-ac3e542f2a8c',status='inspection_in_progress',files=rows),ensure_ascii=False,indent=2),encoding='utf-8')
